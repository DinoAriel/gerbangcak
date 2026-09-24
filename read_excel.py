import openpyxl
import zipfile
import base64
import json
import xml.etree.ElementTree as ET
from openpyxl import load_workbook

EXCEL_FILE = r'c:\kuliah\gerbang\gerbangcak\Rekap All_Data Kendaraan LT 2026_Seleksi.xlsx'

# =============================================
# 1. BACA MAPPING GAMBAR -> SEL DI EXCEL
# =============================================
def get_image_cell_mapping(zip_file, drawing_rel_file, drawing_file):
    """Baca mapping: image file -> anchor cell (baris, kolom)"""
    mapping = {}
    ns = {
        'xdr': 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing',
        'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
        'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    }
    ns_rel = {'': 'http://schemas.openxmlformats.org/package/2006/relationships'}

    # Baca rel file untuk dapat nama file gambar
    rel_map = {}
    if drawing_rel_file in zip_file.namelist():
        with zip_file.open(drawing_rel_file) as f:
            tree = ET.parse(f)
            for rel in tree.getroot():
                rid = rel.get('Id')
                target = rel.get('Target')
                rel_map[rid] = target  # e.g. '../media/image1.png'

    # Baca drawing XML untuk dapat posisi anchor
    if drawing_file in zip_file.namelist():
        with zip_file.open(drawing_file) as f:
            tree = ET.parse(f)
            root = tree.getroot()
            for anchor in root:
                # twoCellAnchor atau oneCellAnchor
                tag = anchor.tag.split('}')[-1] if '}' in anchor.tag else anchor.tag
                if 'Anchor' in tag:
                    # Ambil from cell
                    from_el = anchor.find('{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}from')
                    if from_el is not None:
                        row_el = from_el.find('{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}row')
                        col_el = from_el.find('{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}col')
                        if row_el is not None and col_el is not None:
                            row = int(row_el.text) + 1  # 0-indexed -> 1-indexed
                            col = int(col_el.text) + 1
                            # Cari blip (gambar ID)
                            for pic in anchor.iter('{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}pic'):
                                blip_fill = pic.find('.//{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}blipFill')
                                if blip_fill is None:
                                    blip_fill = pic
                                for blip in pic.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}blip'):
                                    rid = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                                    if rid and rid in rel_map:
                                        img_name = rel_map[rid].replace('../media/', 'xl/media/')
                                        mapping[(row, col)] = img_name
    return mapping

# Load semua gambar dari ZIP
all_images = {}
with zipfile.ZipFile(EXCEL_FILE, 'r') as z:
    for img_path in z.namelist():
        if img_path.startswith('xl/media/'):
            with z.open(img_path) as f:
                all_images[img_path] = base64.b64encode(f.read()).decode('utf-8')

    # Sheet ke drawing file mapping
    sheet_drawing_map = {
        'ANGKUTAN TAKSI': ('xl/drawings/_rels/drawing1.xml.rels', 'xl/drawings/drawing1.xml'),
        'ANGKUTAN SEWA': (None, 'xl/drawings/drawing2.xml'),
        'ANTAR JEMPUT': ('xl/drawings/_rels/drawing3.xml.rels', 'xl/drawings/drawing3.xml'),
    }

    image_mappings = {}
    for sheet_name, (rel_file, draw_file) in sheet_drawing_map.items():
        if rel_file:
            mapping = get_image_cell_mapping(z, rel_file, draw_file)
            image_mappings[sheet_name] = mapping
            print(f"Sheet '{sheet_name}': {len(mapping)} gambar terdeteksi")
            for (r, c), img in mapping.items():
                print(f"   Row {r}, Col {c} -> {img}")
        else:
            image_mappings[sheet_name] = {}
            print(f"Sheet '{sheet_name}': Tidak ada rel file, skip gambar")

# =============================================
# 2. BACA DATA EXCEL + GABUNGKAN DENGAN GAMBAR
# =============================================
wb = load_workbook(EXCEL_FILE, data_only=True)
result = {'kendaraans': [], 'pengemudis': [], 'relations': []}
pengemudi_seen = {}  # nama -> id untuk deduplikasi
pengemudi_counter = [0]

def get_or_create_pengemudi(nama, foto_b64=None, foto_mime=None):
    key = nama.strip().upper()
    if key not in pengemudi_seen:
        pengemudi_counter[0] += 1
        pid = pengemudi_counter[0]
        pengemudi_seen[key] = pid
        result['pengemudis'].append({
            'id': pid,
            'nama': nama.strip(),
            'no_hp': '-',
            'no_sim': '-',
            'status': 'aktif',
            'foto': foto_b64,
            'foto_mime': foto_mime,
        })
    return pengemudi_seen[key]

kendaraan_counter = 0

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    headers = [str(cell.value).strip() if cell.value else '' for cell in ws[1]]
    img_map = image_mappings.get(sheet_name, {})
    
    # Tentukan kolom pengemudi dan foto
    pengemudi_cols = []
    for i, h in enumerate(headers):
        col_num = i + 1
        if 'PENGEMUDI' in h.upper() and 'FOTO' not in h.upper():
            # Cek apakah kolom berikutnya adalah FOTO
            next_h = headers[i+1].upper() if i+1 < len(headers) else ''
            foto_col = (col_num + 1) if 'FOTO' in next_h else None
            pengemudi_cols.append((col_num, foto_col))
    
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        if not any(row[j] for j in range(min(6, len(row)))):
            continue
        
        no = row[0]
        kategori = row[1]
        perusahaan = row[2]
        brand = row[3]
        nomor_kendaraan = row[4]
        tahun = row[5]
        
        if not nomor_kendaraan or not kategori:
            continue
        
        # Bersihkan nomor kendaraan
        nomor_kendaraan = str(nomor_kendaraan).strip()
        tahun_int = int(float(str(tahun))) if tahun else 0
        kode_unik = f"KND-{nomor_kendaraan.replace(' ', '')}"
        
        kendaraan_counter += 1
        kid = kendaraan_counter
        result['kendaraans'].append({
            'id': kid,
            'kode_unik': kode_unik,
            'kategori': str(kategori).strip(),
            'perusahaan': str(perusahaan).strip() if perusahaan else '-',
            'brand': str(brand).strip() if brand else '-',
            'nomor_kendaraan': nomor_kendaraan,
            'tahun_pembuatan': tahun_int,
        })
        
        # Proses pengemudi
        actual_row = row_idx + 2  # 1-indexed, +1 for header
        for (p_col, f_col) in pengemudi_cols:
            p_idx = p_col - 1
            nama_pengemudi = row[p_idx] if p_idx < len(row) else None
            if not nama_pengemudi:
                continue
            nama_pengemudi = str(nama_pengemudi).strip()
            if not nama_pengemudi or nama_pengemudi.upper() == 'NONE':
                continue
            
            # Cek apakah ada foto di sel ini
            foto_b64 = None
            foto_mime = None
            if f_col:
                img_key = (actual_row, f_col)
                if img_key in img_map:
                    img_path = img_map[img_key]
                    if img_path in all_images:
                        foto_b64 = all_images[img_path]
                        ext = img_path.split('.')[-1].lower()
                        foto_mime = f'image/{ext}' if ext != 'jpg' else 'image/jpeg'
            
            pid = get_or_create_pengemudi(nama_pengemudi, foto_b64, foto_mime)
            result['relations'].append({'kendaraan_id': kid, 'pengemudi_id': pid})

print(f"\n=== HASIL EKSTRAKSI ===")
print(f"Total Kendaraan : {len(result['kendaraans'])}")
print(f"Total Pengemudi : {len(result['pengemudis'])}")
print(f"Total Relasi    : {len(result['relations'])}")
print(f"Pengemudi dgn foto: {sum(1 for p in result['pengemudis'] if p['foto'])}")

# Simpan ke JSON (tanpa foto untuk preview)
preview = {
    'kendaraans': result['kendaraans'][:5],
    'pengemudis': [{'id': p['id'], 'nama': p['nama'], 'has_foto': bool(p['foto'])} for p in result['pengemudis'][:20]],
    'relations': result['relations'][:10],
}
with open(r'c:\kuliah\gerbang\gerbangcak\excel_preview.json', 'w', encoding='utf-8') as f:
    json.dump(preview, f, ensure_ascii=False, indent=2)

# Simpan data lengkap ke JSON (dengan foto)
with open(r'c:\kuliah\gerbang\gerbangcak\excel_full_data.json', 'w', encoding='utf-8') as f:
    json.dump(result, f, ensure_ascii=False)

print("\nData tersimpan ke excel_full_data.json dan excel_preview.json")
